import tkinter as tk
import ttkbootstrap as ttk
from tkinter import filedialog
from tkinter import messagebox
import docx
import PyPDF2
import time
import os
import sys
import threading
import re
import csv
import json
import hashlib
import shutil
import base64
import secrets
from pathlib import Path

ROAD_VARIATIONS = {
    "street": ["street", "st", "st."],
    "avenue": ["avenue", "ave", "ave."],
    "road": ["road", "rd", "rd."],
    "boulevard": ["boulevard", "blvd", "blvd."],
    "drive": ["drive", "dr", "dr."],
    "lane": ["lane", "ln", "ln."],
    "court": ["court", "ct", "ct."],
    "terrace": ["terrace", "ter", "ter.", "terr"],
    "place": ["place", "pl", "pl."],
    "square": ["square", "sq", "sq."],
    "highway": ["highway", "hwy", "hwy."],
    "parkway": ["parkway", "pkwy", "pkwy."],
    "circle": ["circle", "cir", "cir."],
    "trail": ["trail", "trl", "trl."],
    "way": ["way", "wy", "wy."]
}

from tts_manager import TTSManager
from text_manager import TextManager
from progress_bar_manager import ProgressBarManager

class AudioTypingTest:
    def __init__(self, root):
        self.root = root
        self.root.title("echoType")

        root.tk.call("source", "azure.tcl")
        root.tk.call("set_theme", "dark")

        self.runtime_dir = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
        self.config_path = self.runtime_dir / "config.json"
        self.app_data_dir = self.load_app_data_dir()
        self.details_dir = self.app_data_dir / "Details"
        self.generations_dir = self.app_data_dir / "Generations"
        self.scores_file = self.app_data_dir / "scores.enc"
        self.tts_temp_file = self.app_data_dir / "TypingTTS.wav"
        self.ensure_app_dirs()
        self.current_detail_key = None
        self.current_file_key = None
        self.current_details = []
        self.details_dialog_open = False
        self.pending_file_loaded_message = False
        self.pending_audio_ready_message = False
        self.speed_dirty = False
        self.generating = False
        self.regeneration_reason = None
        self.current_username = None
        self.current_first_name = None
        self.current_last_name = None
        self.current_is_admin = False
        self._invalid_password_after_id = None
        self._password_fg = None
        self._password_show = None

        self._install_messagebox_topmost()
        self._install_filedialog_topmost()
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        self.root.geometry(f"{screen_width}x{screen_height}")

        self.voice_options = {
            "English": "en_US-libritts-high.onnx",
            "Spanish": "es_MX-claude-high.onnx"
        }
        self.language_var = tk.StringVar(value="English")
        self.current_language = "English"

        self.setup_ui()
        self.tts_manager = TTSManager(filename=str(self.tts_temp_file))
        self.tts_from_file = False
        self.progress_bar_manager = ProgressBarManager(self.root, self.tts_manager)
        self.progress_bar_manager.set_on_complete(self.handle_playback_complete)
        self.text_manager = TextManager(self.root)
        self.text_manager.typing_box.bind("<KeyRelease>", self.on_typing)
        self.text_manager.typing_box.bind("<KeyPress>", self.start_timer_if_needed)
        self.start_time = None
        self.timer_id = None  # For scheduling timer updates
        self.road_variant_map = self.build_road_variant_map()
        self.apply_saved_settings()
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

    # Configuration and path utilities
    def default_app_data_dir(self):
        plat = sys.platform
        if plat.startswith("win"):
            return Path("C:/Program Files/echoType")
        if plat == "darwin":
            return Path("/Applications/echoType")
        return Path("/usr/bin/echoType")

    def load_config(self):
        if self.config_path.exists():
            try:
                with open(self.config_path, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception:
                return {}
        # Auto-create a default config if missing
        default_dir = getattr(self, "app_data_dir", None) or self.default_app_data_dir()
        default = {
            "app_data_dir": str(default_dir),
            "encryption_key": secrets.token_hex(32),
            "ui_settings": self.default_ui_settings()
        }
        try:
            with open(self.config_path, "w", encoding="utf-8") as f:
                json.dump(default, f, ensure_ascii=False, indent=2)
        except Exception:
            pass
        return default

    def save_config(self):
        config = self.load_config()
        config["app_data_dir"] = str(self.app_data_dir)
        if "encryption_key" not in config:
            config["encryption_key"] = secrets.token_hex(32)
        if "ui_settings" not in config:
            config["ui_settings"] = self.get_current_ui_settings()
        try:
            with open(self.config_path, "w", encoding="utf-8") as f:
                json.dump(config, f, ensure_ascii=False, indent=2)
        except Exception as exc:
            messagebox.showerror("Config Error", f"Could not save configuration:\n{exc}")

    def load_app_data_dir(self):
        config = self.load_config()
        configured = config.get("app_data_dir")
        if configured:
            candidate = Path(configured).expanduser()
            if not candidate.is_absolute():
                candidate = (self.runtime_dir / candidate).resolve()
        else:
            candidate = self.default_app_data_dir()
        try:
            candidate.mkdir(parents=True, exist_ok=True)
            return candidate
        except Exception:
            fallback = self.runtime_dir / "echoType_data"
            fallback.mkdir(parents=True, exist_ok=True)
            return fallback

    def ensure_app_dirs(self):
        for path in [self.app_data_dir, self.details_dir, self.generations_dir]:
            path.mkdir(parents=True, exist_ok=True)

    def update_app_data_dir(self, new_dir: Path):
        new_dir = Path(new_dir).expanduser()
        if not new_dir.is_absolute():
            new_dir = (self.runtime_dir / new_dir).resolve()

        if new_dir == self.app_data_dir:
            return

        try:
            new_dir.mkdir(parents=True, exist_ok=True)
        except Exception as exc:
            messagebox.showerror("Path Error", f"Could not use the selected directory:\n{exc}")
            return

        self.app_data_dir = new_dir
        self.details_dir = self.app_data_dir / "Details"
        self.generations_dir = self.app_data_dir / "Generations"
        self.scores_file = self.app_data_dir / "scores.enc"
        self.tts_temp_file = self.app_data_dir / "TypingTTS.wav"
        self.ensure_app_dirs()
        self.tts_manager.filename = str(self.tts_temp_file)
        self.tts_manager.wav_file = str(self.tts_temp_file)
        self.save_config()

    # Simple XOR + base64 helpers for lightweight encryption of stored data
    def _get_encryption_key(self) -> bytes:
        config = self.load_config()
        key_hex = config.get("encryption_key")
        if not key_hex:
            key_hex = secrets.token_hex(32)
            config["encryption_key"] = key_hex
            try:
                with open(self.config_path, "w", encoding="utf-8") as f:
                    json.dump(config, f, ensure_ascii=False, indent=2)
            except Exception:
                pass
        return bytes.fromhex(key_hex)

    def _xor_bytes(self, data: bytes, key: bytes) -> bytes:
        if not key:
            return data
        return bytes(b ^ key[i % len(key)] for i, b in enumerate(data))

    def encrypt_payload(self, payload: dict) -> bytes:
        raw = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        key = self._get_encryption_key()
        cipher = self._xor_bytes(raw, key)
        return base64.b64encode(cipher)

    def decrypt_payload(self, data: bytes) -> dict:
        try:
            key = self._get_encryption_key()
            raw = base64.b64decode(data)
            plain = self._xor_bytes(raw, key)
            return json.loads(plain.decode("utf-8"))
        except Exception:
            return {}

    def default_ui_settings(self):
        return {
            "distortion": "off_distortion",
            "language": "English",
            "speed": 1.0,
            "highlight": "off_highlight"
        }

    def get_current_ui_settings(self):
        try:
            return {
                "distortion": self.distortion_status.get(),
                "language": self.language_var.get(),
                "speed": self.speed_var.get(),
                "highlight": self.highlight_var.get()
            }
        except Exception:
            return self.default_ui_settings()

    def load_ui_settings(self):
        config = self.load_config()
        return config.get("ui_settings", {})

    def save_ui_settings(self):
        if not self.current_is_admin:
            return
        if not hasattr(self, "distortion_status") or not hasattr(self, "language_var"):
            return
        config = self.load_config()
        config["app_data_dir"] = str(self.app_data_dir)
        if "encryption_key" not in config:
            config["encryption_key"] = secrets.token_hex(32)
        config["ui_settings"] = self.get_current_ui_settings()
        try:
            with open(self.config_path, "w", encoding="utf-8") as f:
                json.dump(config, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def fit_window_to_content(self, window, min_size=(400, 300), pad=(40, 40)):
        """Resize a toplevel to fit its content with sensible padding and screen bounds."""
        window.update_idletasks()
        req_w = window.winfo_reqwidth() + pad[0]
        req_h = window.winfo_reqheight() + pad[1]
        screen_w = window.winfo_screenwidth()
        screen_h = window.winfo_screenheight()
        width = max(min_size[0], min(req_w, screen_w - 40))
        height = max(min_size[1], min(req_h, screen_h - 40))
        window.geometry(f"{int(width)}x{int(height)}")

    def bring_window_to_front(self, window):
        """Ensure a toplevel appears above other windows on creation."""
        try:
            window.attributes("-topmost", True)
        except Exception:
            pass
        window.lift()
        window.focus_force()
        # Allow normal stacking after initial display
        window.after(200, lambda: window.attributes("-topmost", False))

    def _install_messagebox_topmost(self):
        """Wrap common messagebox calls so they float above other windows."""
        def make_wrapper(func):
            def wrapper(*args, **kwargs):
                args, kwargs = self._normalize_messagebox_args(*args, **kwargs)
                try:
                    self.root.attributes("-topmost", True)
                except Exception:
                    pass
                try:
                    return func(*args, **kwargs)
                finally:
                    try:
                        self.root.after(200, lambda: self.root.attributes("-topmost", False))
                    except Exception:
                        pass
            return wrapper

        for name in ["showinfo", "showwarning", "showerror", "askyesno", "askyesnocancel", "askquestion"]:
            original = getattr(messagebox, name, None)
            if callable(original):
                setattr(messagebox, name, make_wrapper(original))

    def _trim_title(self, title: str):
        # Remove titles entirely to avoid truncation/ellipsis in native dialogs
        return ""

    def _install_filedialog_topmost(self):
        """Wrap file dialogs so they appear above other windows."""
        def make_wrapper(func):
            def wrapper(*args, **kwargs):
                kwargs.setdefault("parent", self.root)
                try:
                    self.root.attributes("-topmost", True)
                    self.root.lift()
                    self.root.focus_force()
                    self.root.update_idletasks()
                except Exception:
                    pass
                try:
                    return func(*args, **kwargs)
                finally:
                    try:
                        self.root.after(200, lambda: self.root.attributes("-topmost", False))
                    except Exception:
                        pass
            return wrapper

        for name in ["askopenfilename", "asksaveasfilename", "askdirectory"]:
            original = getattr(filedialog, name, None)
            if callable(original):
                setattr(filedialog, name, make_wrapper(original))

    def _normalize_messagebox_args(self, *args, **kwargs):
        args_list = list(args)
        if args_list:
            args_list[0] = self._trim_title(args_list[0])
        elif "title" in kwargs:
            kwargs["title"] = self._trim_title(kwargs.get("title"))
        else:
            args_list.insert(0, self._trim_title(None))
        return args_list, kwargs

    def setup_ui(self):
        self.root.columnconfigure(0, weight=1)
        self.root.columnconfigure(1, weight=3)
        self.root.columnconfigure(2, weight=1)
        self.root.rowconfigure(0, weight=1)
        self.root.rowconfigure(1, weight=4)
        self.root.rowconfigure(2, weight=1)

        tk.Label(self.root, text="echoType", font=("Calibri", 16, "bold")).grid(row=0, column=1, pady=10, sticky="n")

        self.sidebar = tk.Frame(self.root, width=500, bd=1 , relief="raised")
        self.sidebar.grid(row=0, column=0, rowspan=3, sticky="nsw")

        self.settings_label = tk.Label(self.sidebar, text='Settings', font=("Calibri", 14))
        self.settings_label.grid(row=0, column=0, padx=10, pady=10, sticky="new")

        self.config_button = tk.Button(self.sidebar, text="Configuration Settings", command=self.open_config_settings)
        self.config_button.grid(row=1, column=0, padx=10, pady=(0, 10), sticky="ew")

        self.distortion_label = tk.Label(self.sidebar, text='Distortion:', font=("Calibri", 12))
        self.distortion_label.grid(row=2, column=0, padx=10, pady=10, sticky="w")

        self.distortion_status = tk.StringVar()
        self.distortion_on = ttk.Radiobutton(
            self.sidebar,
            text="On",
            variable=self.distortion_status,
            value="on_distortion",
            command=self.update_distortion_setting
        )
        self.distortion_on.grid(row=3, column=0, padx=10, pady=5, sticky="ew")
        self.distortion_off = ttk.Radiobutton(
            self.sidebar,
            text="Off",
            variable=self.distortion_status,
            value="off_distortion",
            command=self.update_distortion_setting
        )

        self.sign_in_button = ttk.Button(self.sidebar, text="Sign In", command=self.sign_in)
        self.sign_in_button.grid(row=12, column=0, padx=10, pady=(15, 5), sticky="ew")
        
        self.register_button = tk.Button(self.sidebar, text="Register New Account", command=lambda: self.open_register_dialog())
        self.register_button.grid(row=13, column=0, padx=10, pady=5, sticky="ew") 

        self.load_file_button = tk.Button(self.sidebar, text="Load File for TTS", command=self.load_file_for_tts)
        self.load_file_button.grid(row=15, column=0, padx=10, pady=10, sticky="ew")

        self.distortion_off.grid(row=4, column=0, padx=10, pady=5, sticky="ew")
        self.distortion_status.set("off_distortion")

        self.language_label = tk.Label(self.sidebar, text="Language:", font=("Calibri", 12))
        self.language_label.grid(row=6, column=0, padx=10, pady=(10, 0), sticky="w")

        self.language_frame = tk.Frame(self.sidebar)
        self.language_frame.grid(row=7, column=0, padx=10, pady=5, sticky="w")

        self.language_en = ttk.Radiobutton(
            self.language_frame,
            text="English",
            variable=self.language_var,
            value="English",
            command=self.change_language
        )
        self.language_en.pack(side="left", padx=(0, 10))

        self.language_es = ttk.Radiobutton(
            self.language_frame,
            text="Spanish",
            variable=self.language_var,
            value="Spanish",
            command=self.change_language
        )
        self.language_es.pack(side="left")

        

        self.username_label = tk.Label(self.sidebar, text="Username:")
        self.username_label.grid(row=8, column=0, padx=10, pady=10, sticky="esw")
        self.username_value = tk.StringVar()
        self.username_entry = tk.Entry(self.sidebar, textvariable=self.username_value)
        self.username_entry.grid(row=9, column=0, padx=10, pady=5, sticky="esw")

        self.password_label = tk.Label(self.sidebar, text="Password:")
        self.password_label.grid(row=10, column=0, padx=10, pady=10, sticky="esw")
        self.password_value = tk.StringVar()
        self.password_entry = tk.Entry(self.sidebar, textvariable=self.password_value, show="*")
        self.password_entry.grid(row=11, column=0, padx=10, pady=5, sticky="esw")

        self.submit_button = tk.Button(self.root, text="Submit", command=self.submit_text)
        self.submit_button.grid(row=2, column=2, padx=10, pady=10, sticky="w")

        self.discard_button = tk.Button(self.root, text="Discard", command=self.discard_text)
        self.discard_button.grid(row=2, column=2, padx=10, pady=10, sticky="e")

        self.play_pause_button = tk.Button(self.root, text="▶", font=("Arial", 14), command=self.toggle_play_pause)
        self.play_pause_button.grid(row=0, column=2, padx=(10, 5), pady=10, sticky="w")

        self.reset_button = tk.Button(self.root, text="⟲", font=("Arial", 14), command=self.reset_audio)
        self.reset_button.grid(row=0, column=2, padx=(5, 10), pady=10, sticky="e")

        self.speed_label = tk.Label(self.sidebar, text='TTS Speed:', font=("Calibri", 12))
        self.speed_label.grid(row=16, column=0, padx=10, pady=(20, 0), sticky="w")

        self.speed_var = tk.DoubleVar(value=1.0)  # Default speed = 1.0
        self.speed_slider = tk.Scale(self.sidebar, from_=0.5, to=2.0, resolution=0.1,
                                    orient="horizontal", variable=self.speed_var,
                                    length=200, command=self.on_speed_dirty)
        self.speed_slider.grid(row=17, column=0, padx=10, pady=5, sticky="ew")
        self.apply_speed_button = tk.Button(self.sidebar, text="Apply Speed (1.0x)", command=self.apply_speed_change, state="disabled")
        self.apply_speed_button.grid(row=18, column=0, padx=10, pady=(5, 15), sticky="ew")

        self.highlight_label = tk.Label(self.sidebar, text="Show Spelling Errors:", font=("Calibri", 12))
        self.highlight_label.grid(row=19, column=0, padx=10, pady=(20, 0), sticky="w")

        self.highlight_var = tk.StringVar(value="off_highlight")  # Default = OFF

        self.highlight_on = ttk.Radiobutton(self.sidebar, text="Yes", variable=self.highlight_var, value="on_highlight", command=self.on_highlight_changed)
        self.highlight_on.grid(row=20, column=0, padx=10, sticky="w")

        self.highlight_off = ttk.Radiobutton(self.sidebar, text="No", variable=self.highlight_var, value="off_highlight", command=self.on_highlight_changed)
        self.highlight_off.grid(row=21, column=0, padx=10, sticky="w")

        self.view_scores_button = tk.Button(self.sidebar, text="View Scores", command=self.open_scores_view)
        self.view_scores_button.grid(row=22, column=0, padx=10, pady=(20, 15), sticky="ew")
        self.update_admin_controls()
        self.update_apply_speed_button()

    def on_close(self):
        # Stop timers/UI loops
        self.stop_timer_display()
        self.progress_bar_manager.reset_progress_bar()

        # Stop audio stream cleanly
        try:
            self.tts_manager.pauseTTS()
            if self.tts_manager.stream:
                try:
                    self.tts_manager.stream.stop()
                except Exception:
                    pass
                try:
                    self.tts_manager.stream.close()
                except Exception:
                    pass
        except Exception:
            pass

        # Delete synthesized files
        try:
            self.tts_manager.deleteTTSFile()
        except Exception:
            pass

        # Close the window
        self.root.destroy()

    # Account Management Configuration
    def get_user_db_path(self):
        return self.details_dir / "users.json"

    def load_user_db(self):
        #Loads the user database
        db_path = self.get_user_db_path()
        if db_path.exists():
            try:
                with open(db_path, 'r', encoding="utf-8") as f:
                    data = json.load(f)
                ensured = self.ensure_admin_seed(data)
                return ensured
            except Exception as exc:
                messagebox.showerror(
                    "User Database Error",
                    f"Could not read the user database at:\n{db_path}\n\nDetails: {exc}\n"
                    "Please fix or replace the file before signing in or registering."
                )
                return None
        return {}

    def save_user_db(self, user_db):
        db_path = self.get_user_db_path()
        self.details_dir.mkdir(parents=True, exist_ok=True)
        with open(db_path, 'w', encoding="utf-8") as f:
            json.dump(user_db, f, indent=4, ensure_ascii=False)

    def hash_password(self, password):
        return hashlib.sha256(password.encode('utf-8')).hexdigest()

    def ensure_admin_seed(self, user_db: dict):
        if not isinstance(user_db, dict):
            return {}
        has_admin = False
        for record in user_db.values():
            if isinstance(record, dict) and record.get("is_admin"):
                has_admin = True
                break
        if has_admin or not user_db:
            return user_db
        # No admin found; promote the first listed user to admin
        first_username = next(iter(user_db.keys()))
        record = user_db[first_username]
        if isinstance(record, dict):
            record["is_admin"] = True
        else:
            user_db[first_username] = {"password_hash": record, "is_admin": True}
        try:
            self.save_user_db(user_db)
        except Exception:
            pass
        return user_db

    def sign_in(self):
        username = self.username_value.get().strip()
        password = self.password_value.get()
        user_db = self.load_user_db()
        if user_db is None:
            return

        if not username or not password:
            messagebox.showwarning("Login Error", "Please enter both username and password.")
            return

        if username in user_db:
            record = user_db[username]
            password_hash = None
            first_name = None
            last_name = None
            is_admin = False
            if isinstance(record, dict):
                password_hash = record.get("password_hash")
                first_name = record.get("first_name")
                last_name = record.get("last_name")
                is_admin = bool(record.get("is_admin"))
            else:
                password_hash = record
            if password_hash and self.hash_password(password) == password_hash:
                messagebox.showinfo("Success", f"Welcome, {username}!")
                self.set_logged_in_state(username, first_name, last_name, is_admin=is_admin)
            else:
                self.show_invalid_password_message()
        else:
            self.handle_invalid_credentials(username, "Username not found.")

    def handle_invalid_credentials(self, username, error_message):
        self.set_logged_in_state(None)
        should_register = messagebox.askyesno(
            "Login Failed",
            f"{error_message} Would you like to register a new account for '{username}'?"
        )
        if should_register:
            self.open_register_dialog(username)

    def sign_out(self):
        """Sign the current user out and re-enable the login form."""
        if not self.current_username:
            return
        self.set_logged_in_state(None)
        messagebox.showinfo("Signed Out", "You have been signed out.")

    def open_config_settings(self):
        if not self.current_is_admin:
            messagebox.showwarning("Admin Only", "Configuration Settings are available to admins only.")
            return
        dialog = tk.Toplevel(self.root)
        dialog.title("Configuration Settings")
        dialog.resizable(True, True)
        dialog.grab_set()
        dialog.transient(self.root)
        self.bring_window_to_front(dialog)

        user_db = self.load_user_db() or {}
        usernames = sorted(user_db.keys())

        # Application data directory section
        data_frame = tk.LabelFrame(dialog, text="Application Data Directory")
        data_frame.pack(fill="x", padx=10, pady=10)

        dir_var = tk.StringVar(value=str(self.app_data_dir))

        tk.Label(data_frame, text="Path:").grid(row=0, column=0, padx=10, pady=5, sticky="e")
        dir_entry = tk.Entry(data_frame, textvariable=dir_var, width=50)
        dir_entry.grid(row=0, column=1, padx=10, pady=5, sticky="we")
        data_frame.columnconfigure(1, weight=1)

        def browse_dir():
            selected = filedialog.askdirectory(initialdir=dir_var.get() or str(self.runtime_dir))
            if selected:
                dir_var.set(selected)

        browse_btn = tk.Button(data_frame, text="Browse...", command=browse_dir)
        browse_btn.grid(row=0, column=2, padx=10, pady=5)

        def save_dir():
            chosen = dir_var.get().strip()
            if not chosen:
                messagebox.showwarning("Invalid Path", "Please choose a valid directory.")
                return
            self.update_app_data_dir(Path(chosen))
            dir_var.set(str(self.app_data_dir))
            messagebox.showinfo("Data Directory Updated", f"Application data directory set to:\n{self.app_data_dir}")

        tk.Button(data_frame, text="Save Directory", command=save_dir).grid(row=1, column=1, padx=10, pady=(0, 10), sticky="e")

        # Admin management section
        admin_frame = tk.LabelFrame(dialog, text="Admin Management")
        admin_frame.pack(fill="x", padx=10, pady=10)

        admin_user_var = tk.StringVar()
        admin_user_menu = ttk.Combobox(admin_frame, textvariable=admin_user_var, values=usernames, state="readonly")
        admin_user_menu.grid(row=0, column=0, padx=10, pady=5, sticky="we")
        admin_frame.columnconfigure(0, weight=1)

        def grant_admin():
            target = admin_user_var.get()
            if not target:
                messagebox.showwarning("Select User", "Select a user to promote to admin.")
                return
            if target == self.current_username:
                messagebox.showinfo("Already Admin", "You are already an admin.")
                return
            db = self.load_user_db() or {}
            rec = db.get(target)
            if isinstance(rec, dict):
                rec["is_admin"] = True
            else:
                db[target] = {"password_hash": rec, "is_admin": True}
            self.save_user_db(db)
            messagebox.showinfo("Admin Granted", f"{target} is now an admin.")
        tk.Button(admin_frame, text="Grant Admin", command=grant_admin).grid(row=0, column=1, padx=10, pady=5, sticky="e")

        # Reset password section
        reset_frame = tk.LabelFrame(dialog, text="Reset User Password (requires current user password)")
        reset_frame.pack(fill="x", padx=10, pady=10)

        reset_user_var = tk.StringVar()
        reset_user_menu = ttk.Combobox(reset_frame, textvariable=reset_user_var, values=usernames, state="readonly")
        reset_user_menu.grid(row=0, column=0, padx=10, pady=5, sticky="we")
        reset_frame.columnconfigure(0, weight=1)

        new_pwd_var = tk.StringVar()
        confirm_pwd_var = tk.StringVar()
        admin_pwd_var = tk.StringVar()

        tk.Label(reset_frame, text="New Password:").grid(row=1, column=0, padx=10, pady=5, sticky="w")
        tk.Entry(reset_frame, textvariable=new_pwd_var, show="*").grid(row=1, column=1, padx=10, pady=5, sticky="we")
        tk.Label(reset_frame, text="Confirm Password:").grid(row=2, column=0, padx=10, pady=5, sticky="w")
        tk.Entry(reset_frame, textvariable=confirm_pwd_var, show="*").grid(row=2, column=1, padx=10, pady=5, sticky="we")
        tk.Label(reset_frame, text="Current Admin Password:").grid(row=3, column=0, padx=10, pady=5, sticky="w")
        tk.Entry(reset_frame, textvariable=admin_pwd_var, show="*").grid(row=3, column=1, padx=10, pady=5, sticky="we")
        reset_frame.columnconfigure(1, weight=1)

        def handle_reset_password():
            target = reset_user_var.get()
            new_pwd = new_pwd_var.get()
            confirm_pwd = confirm_pwd_var.get()
            admin_pwd = admin_pwd_var.get()

            if not target:
                messagebox.showwarning("Select User", "Select a user to reset.")
                return
            if not new_pwd or not confirm_pwd or not admin_pwd:
                messagebox.showwarning("Missing Information", "Enter new password, confirm it, and provide your admin password.")
                return
            if new_pwd != confirm_pwd:
                messagebox.showwarning("Password Mismatch", "New password and confirmation do not match.")
                return
            if len(new_pwd) < 4:
                messagebox.showwarning("Weak Password", "Password must be at least 4 characters.")
                return
            db = self.load_user_db() or {}
            if target not in db:
                messagebox.showwarning("User Missing", f"User '{target}' does not exist.")
                return
            if not self.current_username:
                messagebox.showwarning("Not Signed In", "Sign in as an admin to reset passwords.")
                return
            current_record = db.get(self.current_username)
            current_hash = current_record.get("password_hash") if isinstance(current_record, dict) else current_record
            if not current_hash or self.hash_password(admin_pwd) != current_hash:
                messagebox.showerror("Auth Failed", "Incorrect admin password.")
                return

            record = db.get(target, {})
            if isinstance(record, dict):
                record["password_hash"] = self.hash_password(new_pwd)
            else:
                record = {"password_hash": self.hash_password(new_pwd)}
            db[target] = record
            self.save_user_db(db)
            new_pwd_var.set("")
            confirm_pwd_var.set("")
            admin_pwd_var.set("")
            messagebox.showinfo("Password Reset", f"Password reset for {target}.")

        tk.Button(reset_frame, text="Reset Password", command=handle_reset_password).grid(row=4, column=1, padx=10, pady=(5, 10), sticky="e")

        # User deletion section
        user_frame = tk.LabelFrame(dialog, text="Delete User and Data (requires current user password)")
        user_frame.pack(fill="both", padx=10, pady=10, expand=True)

        tk.Label(user_frame, text="Select User:").grid(row=0, column=0, padx=10, pady=5, sticky="w")
        user_listbox = tk.Listbox(user_frame, height=8, exportselection=False)
        user_listbox.grid(row=1, column=0, padx=10, pady=5, sticky="nsew")
        for name in usernames:
            user_listbox.insert("end", name)

        pwd_var = tk.StringVar()
        tk.Label(user_frame, text="Current User Password:").grid(row=0, column=1, padx=10, pady=5, sticky="w")
        tk.Entry(user_frame, textvariable=pwd_var, show="*").grid(row=1, column=1, padx=10, pady=5, sticky="we")
        user_frame.columnconfigure(0, weight=1)
        user_frame.columnconfigure(1, weight=1)

        def refresh_user_list():
            user_listbox.delete(0, "end")
            refreshed_db = self.load_user_db() or {}
            for name in sorted(refreshed_db.keys()):
                user_listbox.insert("end", name)
            reset_user_menu["values"] = sorted((self.load_user_db() or {}).keys())

        def selected_user():
            selection = user_listbox.curselection()
            if selection:
                return selection[0], user_listbox.get(selection[0])
            return None, None

        def show_deleted_state(index):
            user_listbox.delete(index)
            user_listbox.insert(index, "deleted")
            try:
                user_listbox.itemconfig(index, {"fg": "red"})
            except Exception:
                pass
            self.root.after(3000, refresh_user_list)

        def handle_delete_user():
            idx, username = selected_user()
            password = pwd_var.get()
            if not username:
                messagebox.showwarning("Select User", "Please select a user to delete.")
                return
            if not self.current_username:
                messagebox.showwarning("Not Signed In", "Please sign in to delete users.")
                return
            if not password:
                messagebox.showwarning("Missing Password", "Enter the current user's password to confirm deletion.")
                return
            if self.delete_user_and_data(username, password):
                pwd_var.set("")
                if idx is not None:
                    show_deleted_state(idx)

        tk.Button(user_frame, text="Delete User", command=handle_delete_user).grid(row=2, column=0, columnspan=2, padx=10, pady=(5, 10))

        # Delete all data section
        all_frame = tk.LabelFrame(dialog, text="Delete All Application Data (requires current user password)")
        all_frame.pack(fill="x", padx=10, pady=10)

        confirm_pwd = tk.StringVar()

        tk.Label(all_frame, text="Current User:").grid(row=0, column=0, padx=10, pady=5, sticky="w")
        current_user_label = tk.Label(all_frame, text=self.current_username or "Not signed in")
        current_user_label.grid(row=0, column=1, padx=10, pady=5, sticky="w")
        tk.Label(all_frame, text="Password:").grid(row=1, column=0, padx=10, pady=5, sticky="w")
        tk.Entry(all_frame, textvariable=confirm_pwd, show="*").grid(row=1, column=1, padx=10, pady=5, sticky="we")
        all_frame.columnconfigure(1, weight=1)

        def handle_delete_all():
            password = confirm_pwd.get()
            if not self.current_username:
                messagebox.showwarning("Not Signed In", "Please sign in to delete all data.")
                return
            if not password:
                messagebox.showwarning("Missing Password", "Enter the current user's password to confirm deletion.")
                return
            if self.delete_all_app_data(self.current_username, password):
                confirm_pwd.set("")
                refresh_user_list()
                current_user_label.config(text=self.current_username or "Not signed in")

        tk.Button(all_frame, text="Delete All Data", command=handle_delete_all).grid(row=2, column=0, columnspan=2, padx=10, pady=(5, 10))
        self.fit_window_to_content(dialog, min_size=(720, 650))

    def open_scores_view(self):
        if not self.current_is_admin:
            messagebox.showwarning("Admin Only", "Scores are available to admins only.")
            return
        dialog = tk.Toplevel(self.root)
        dialog.title("View Scores")
        dialog.resizable(True, True)
        dialog.grab_set()
        dialog.transient(self.root)
        self.bring_window_to_front(dialog)

        data = self.load_scores_store() or {}
        user_db = self.load_user_db() or {}
        users = sorted(data.keys())

        top_frame = tk.Frame(dialog)
        top_frame.pack(fill="x", padx=10, pady=10)

        tk.Label(top_frame, text="Users:").pack(side="left")
        user_var = tk.StringVar(value=users[0] if users else "")
        user_menu = ttk.Combobox(top_frame, textvariable=user_var, values=users, state="readonly")
        user_menu.pack(side="left", padx=10)

        user_label = tk.Label(top_frame, text=f"User: {user_var.get() or 'None'}")
        user_label.pack(side="left", padx=10)

        name_label = tk.Label(top_frame, text="Name: ")
        name_label.pack(side="left", padx=10)

        table_frame = tk.Frame(dialog)
        table_frame.pack(fill="both", expand=True, padx=10, pady=10)

        columns = ("test_no", "time", "wpm", "accuracy", "details")
        tree = ttk.Treeview(table_frame, columns=columns, show="headings")
        headers = {
            "test_no": "Test No.",
            "time": "Time",
            "wpm": "WPM",
            "accuracy": "Accuracy (%)",
            "details": "Details (%)"
        }
        for col, text in headers.items():
            tree.heading(col, text=text)
            tree.column(col, width=120, anchor="center")

        tree.pack(fill="both", expand=True, side="left")
        scrollbar = ttk.Scrollbar(table_frame, orient="vertical", command=tree.yview)
        scrollbar.pack(side="right", fill="y")
        tree.configure(yscrollcommand=scrollbar.set)

        def populate(user):
            user_label.config(text=f"User: {user or 'None'}")
            tree.delete(*tree.get_children())
            records = data.get(user, [])
            for rec in records:
                tree.insert("", "end", values=(
                    rec.get("test_no", ""),
                    rec.get("time", ""),
                    rec.get("wpm", ""),
                    rec.get("accuracy", ""),
                    rec.get("details", "")
                ))
            record = user_db.get(user, {}) if isinstance(user_db.get(user, {}), dict) else {}
            fname = record.get("first_name") or (records[0].get("first_name") if records else "N/A")
            lname = record.get("last_name") or (records[0].get("last_name") if records else "N/A")
            name_label.config(text=f"Name: {fname} {lname}")

        def on_select(event=None):
            current = user_var.get()
            if current:
                populate(current)

        user_menu.bind("<<ComboboxSelected>>", on_select)
        if users:
            populate(users[0])

        def download_csv():
            current = user_var.get()
            if not current:
                messagebox.showwarning("No User Selected", "Please select a user to download scores.")
                return
            records = data.get(current, [])
            if not records:
                messagebox.showinfo("No Scores", "No scores available for this user.")
                return
            path = filedialog.asksaveasfilename(
                defaultextension=".csv",
                filetypes=[("CSV Files", "*.csv"), ("All Files", "*.*")],
                initialfile=f"{current}_scores.csv"
            )
            if not path:
                return
            header = ["Test No.", "Time", "WPM", "Accuracy (%)", "Details (%)"]
            try:
                with open(path, "w", newline="", encoding="utf-8") as f:
                    writer = csv.writer(f)
                    writer.writerow(header)
                    for rec in records:
                        writer.writerow([
                            rec.get("test_no", ""),
                            rec.get("time", ""),
                            rec.get("wpm", ""),
                            rec.get("accuracy", ""),
                            rec.get("details", "")
                        ])
                messagebox.showinfo("Download Complete", f"Scores saved to {path}")
            except Exception as exc:
                messagebox.showerror("Download Error", f"Could not save scores:\n{exc}")

        download_btn = tk.Button(dialog, text="Download Selected User Scores", command=download_csv)
        download_btn.pack(pady=(0, 10))
        self.fit_window_to_content(dialog, min_size=(820, 520))

    def open_register_dialog(self, suggested_username=None):
        dialog = tk.Toplevel(self.root)
        dialog.title("Register New Account")
        dialog.resizable(True, True)
        dialog.grab_set()
        dialog.transient(self.root)
        self.bring_window_to_front(dialog)

        first_name_var = tk.StringVar()
        last_name_var = tk.StringVar()
        username_var = tk.StringVar(value=suggested_username or "")
        password_var = tk.StringVar()
        confirm_var = tk.StringVar()

        labels = ["First Name", "Last Name", "Username", "Password", "Confirm Password"]
        vars = [first_name_var, last_name_var, username_var, password_var, confirm_var]
        for idx, (label_text, var) in enumerate(zip(labels, vars)):
            tk.Label(dialog, text=label_text + ":").grid(row=idx, column=0, padx=10, pady=5, sticky="e")
            show = "*" if "Password" in label_text else None
            entry = tk.Entry(dialog, textvariable=var, show=show)
            entry.grid(row=idx, column=1, padx=10, pady=5, sticky="we")

        dialog.columnconfigure(1, weight=1)

        def submit():
            if self.register_account(
                first_name_var.get().strip(),
                last_name_var.get().strip(),
                username_var.get().strip(),
                password_var.get(),
                confirm_var.get()
            ):
                dialog.destroy()

        def cancel():
            dialog.destroy()

        button_frame = tk.Frame(dialog)
        button_frame.grid(row=len(labels), column=0, columnspan=2, pady=15)

        tk.Button(button_frame, text="Cancel", command=cancel).pack(side="right", padx=5)
        tk.Button(button_frame, text="Register", command=submit).pack(side="right", padx=5)

        self.fit_window_to_content(dialog, min_size=(400, 300))
        dialog.wait_window()

    def register_account(self, first_name, last_name, username, password, confirm_password):
        user_db = self.load_user_db()
        if user_db is None:
            return False

        if not first_name or not last_name or not username or not password or not confirm_password:
            messagebox.showerror("Registration Error", "All fields are required.")
            return False
        if password != confirm_password:
            messagebox.showerror("Registration Error", "Password and Confirm Password must match.")
            return False
        if len(password) < 4:
            messagebox.showerror("Registration Error", "Password must be at least 4 characters long.")
            return False

        if username in user_db:
            messagebox.showwarning("Registration Error", f"Account for '{username}' already exists. Please sign in.")
            return False

        # Register the new user
        is_first_user = len(user_db) == 0
        user_db[username] = {
            "password_hash": self.hash_password(password),
            "first_name": first_name,
            "last_name": last_name,
            "is_admin": is_first_user
        }
        self.save_user_db(user_db)
        
        messagebox.showinfo("Registration Successful", f"Account created for {username}. You are now signed in.")
        self.set_logged_in_state(username, first_name, last_name, is_admin=is_first_user)
        return True

    def set_logged_in_state(self, username, first_name=None, last_name=None, is_admin=False):
        """Updates the UI based on logged-in status."""
        
        if username:
            self.sign_in_button.config(
                text=f"Sign Out ({username})",
                state="normal",
                style="success.TButton",
                command=self.sign_out
            )
            self.username_entry.config(state="disabled")
            self.password_entry.config(state="disabled")
            # Clear password field after successful login
            self.password_value.set("") 
            self.username_value.set(username)
        else:
            self.sign_in_button.config(
                text="Sign In",
                state="normal",
                style="TButton",
                command=self.sign_in
            )
            self.username_entry.config(state="normal")
            self.password_entry.config(state="normal")
            self.password_value.set("")
            self.username_value.set("")

        self.current_username = username
        self.current_first_name = first_name if username else None
        self.current_last_name = last_name if username else None
        self.current_is_admin = bool(is_admin and username)
        self.update_admin_controls()
        self.save_ui_settings()

    def update_admin_controls(self):
        # Enable/disable admin-only controls
        admin_widgets = [
            self.config_button,
            self.distortion_on,
            self.distortion_off,
            getattr(self, "language_en", None),
            getattr(self, "language_es", None),
            self.speed_slider,
            self.apply_speed_button,
            self.highlight_on,
            self.highlight_off,
            self.view_scores_button
        ]
        state = "normal" if self.current_is_admin else "disabled"
        for widget in admin_widgets:
            if widget:
                widget.config(state=state)

    def delete_user_and_data(self, username, current_user_password):
        user_db = self.load_user_db()
        if user_db is None:
            return False
        if username not in user_db:
            messagebox.showerror("Delete User", f"User '{username}' does not exist.")
            return False
        if not self.current_username:
            messagebox.showerror("Delete User", "You must be signed in to delete a user.")
            return False

        current_record = user_db.get(self.current_username)
        current_hash = current_record.get("password_hash") if isinstance(current_record, dict) else current_record
        if not current_hash or self.hash_password(current_user_password) != current_hash:
            messagebox.showerror("Delete User", "Incorrect current user password. User was not deleted.")
            return False

        user_db.pop(username, None)
        self.save_user_db(user_db)
        self.remove_scores_for_user(username)

        if self.current_username == username:
            self.set_logged_in_state(None)

        return True

    def remove_scores_for_user(self, username):
        path = self.scores_file
        if not path.exists():
            return
        try:
            data = self.load_scores_store() or {}
            if username in data:
                data.pop(username, None)
                self.save_scores_store(data)
        except Exception:
            return

    def delete_all_app_data(self, username, password):
        user_db = self.load_user_db()
        if user_db is None or username not in user_db or not self.current_username:
            messagebox.showerror("Delete Data", "You must be signed in to delete data.")
            return False

        record = user_db[self.current_username]
        stored_hash = record.get("password_hash") if isinstance(record, dict) else record
        if not stored_hash or self.hash_password(password) != stored_hash:
            messagebox.showerror("Delete Data", "Incorrect password. No data was deleted.")
            return False

        errors = []

        for path in [self.details_dir, self.generations_dir]:
            try:
                if path.exists():
                    shutil.rmtree(path, ignore_errors=False)
            except Exception as exc:
                errors.append(str(exc))

        for file_path in [self.scores_file, self.tts_temp_file, self.config_path]:
            try:
                if Path(file_path).exists():
                    Path(file_path).unlink()
            except Exception as exc:
                errors.append(str(exc))

        # Re-create clean directories and empty databases
        self.ensure_app_dirs()
        try:
            self.save_user_db({})
        except Exception as exc:
            errors.append(str(exc))

        self.current_details = []
        self.current_detail_key = None
        self.current_file_key = None
        self.stop_timer_display()
        self.tts_manager.deleteTTSFile()
        self.set_logged_in_state(None)
        self.text_manager.clear_text()
        self.text_manager.hide_results()
        self.progress_bar_manager.reset_progress_bar()

        if errors:
            messagebox.showerror("Delete Data", "Some data could not be deleted:\n" + "\n".join(errors))
            return False

        messagebox.showinfo("Delete Data", "All application data has been deleted.")
        return True

    def show_invalid_password_message(self):
        if self._invalid_password_after_id:
            try:
                self.root.after_cancel(self._invalid_password_after_id)
            except Exception:
                pass
            self._invalid_password_after_id = None

        if self._password_fg is None:
            self._password_fg = self.password_entry.cget("fg")
        if self._password_show is None:
            self._password_show = self.password_entry.cget("show")

        self.password_entry.config(fg="red", show="")
        self.password_value.set("Invalid Password!")

        def reset_password_field():
            self.password_value.set("")
            self.password_entry.config(fg=self._password_fg, show=self._password_show)
            self._invalid_password_after_id = None

        self._invalid_password_after_id = self.root.after(3000, reset_password_field)

    def load_file_for_tts(self):
        file_path = filedialog.askopenfilename(
            title="Select a Text Document",
            filetypes=[
                ("All Supported", "*.txt *.docx *.pdf"),
                ("Text Files", "*.txt"),
                ("Word Documents", "*.docx"),
                ("PDF Files", "*.pdf")
            ]
        )

        if not file_path:
            return

        text_content = ""

        try:
            if file_path.endswith(".txt"):
                with open(file_path, 'r', encoding='utf-8') as file:
                    text_content = file.read()
            elif file_path.endswith(".docx"):
                doc = docx.Document(file_path)
                text_content = "\n".join(p.text for p in doc.paragraphs)
            elif file_path.endswith(".pdf"):
                with open(file_path, 'rb') as file:
                    reader = PyPDF2.PdfReader(file)
                    text_content = "\n".join(p.extract_text() for p in reader.pages if p.extract_text())

            self.tts_manager.typingText = text_content
            self.tts_from_file = True
            

            self.start_time = None

            file_key = self.get_file_key(file_path)
            self.current_file_key = file_key
            generation_path = self.get_generation_path(file_key)

            if generation_path.is_file():
                reuse_audio = messagebox.askyesno(
                    "Existing Audio Found",
                    f"Audio for this document already exists in {self.current_language}.\n"
                    "Select Yes to reuse it or No to regenerate."
                )
                if reuse_audio and self.load_existing_generation(generation_path, text_content):
                    self.handle_details_for_file(file_key, text_content)
                    return

            self.generate_tts_in_background(text_content, save_key=file_key, language=self.current_language)
            self.handle_details_for_file(file_key, text_content)

        except Exception as e:
            messagebox.showerror("Error", f"Could not load file:\n{str(e)}")

    def get_file_key(self, file_path):
        abs_path = os.path.abspath(file_path)
        return hashlib.sha256(abs_path.encode("utf-8")).hexdigest()

    def get_language_slug(self, language=None):
        language = language or self.current_language
        return "es" if language == "Spanish" else "en"

    def get_generation_path(self, file_key, language=None):
        slug = self.get_language_slug(language)
        slugged = self.generations_dir / f"{file_key}_{slug}.wav"
        legacy = self.generations_dir / f"{file_key}.wav"
        if language is not None:
            if slugged.exists():
                return slugged
            if legacy.exists() and (language == "English" or language is None):
                return legacy
            return slugged
        if slugged.exists():
            return slugged
        if legacy.exists():
            return legacy
        return slugged

    def load_existing_generation(self, generation_path, text_content, show_message=True):
        try:
            shutil.copy2(generation_path, self.tts_manager.wav_file)
            self.tts_manager.typingText = text_content
            self.tts_manager._last_text = text_content
            target_scale = self.tts_manager._to_piper_scale(self.speed_var.get())
            self.tts_manager._last_synth_scale = target_scale
            self.tts_manager.load_audio()
            self.tts_manager.is_armed = True
            self.tts_manager.is_paused = True
            self.progress_bar_manager.update_audio_duration(speed=self.speed_var.get())
            self.progress_bar_manager.reset_progress_bar()
            self.update_play_pause_button(False)
            self.speed_dirty = False
            self.update_apply_speed_button()
            if show_message:
                messagebox.showinfo("Audio Loaded", "Existing audio has been loaded for this document.")
            return True
        except Exception as exc:
            messagebox.showerror("Audio Error", f"Failed to load saved audio:\n{exc}\nA new version will be generated.")
            return False

    def handle_details_for_file(self, file_key, text_content):
        self.current_detail_key = file_key
        details_path = self.details_dir / f"{file_key}.json"
        saved_details = self.load_saved_details(details_path)

        if not self.current_is_admin:
            # Non-admins simply reuse saved details if present; otherwise none
            self.current_details = saved_details or []
            return

        if saved_details:
            reuse_saved = messagebox.askyesno(
                "Saved Details Detected",
                "Previously saved details were found for this file.\n"
                "Select Yes to reuse them or No to choose new details."
            )
            if reuse_saved:
                self.current_details = saved_details
                return

        selection = self.show_details_selection_dialog(text_content, saved_details or [])
        if selection is None:
            self.current_details = saved_details or []
            return

        self.current_details = selection
        self.save_details(details_path, selection)

    def update_distortion_setting(self, force=False):
        if not force and not self.current_is_admin:
            return
        enabled = self.distortion_status.get() == "on_distortion"
        if hasattr(self, "tts_manager"):
            self.tts_manager.set_distortion_enabled(enabled)
        if self.current_is_admin:
            self.save_ui_settings()

    def change_language(self, force=False):
        if not force and not self.current_is_admin:
            messagebox.showwarning("Admin Only", "Language changes are available to admins only.")
            return
        selection = self.language_var.get()
        previous_language = getattr(self, "current_language", "English")
        if selection == previous_language:
            return

        model_name = self.voice_options.get(selection)
        if not model_name or not hasattr(self, "tts_manager"):
            self.language_var.set(previous_language)
            return

        def revert_selection():
            self.language_var.set(previous_language)

        try:
            existing_text = self.tts_manager.getTypingText()
        except Exception:
            existing_text = ""
        has_audio_text = bool(existing_text.strip())

        # If nothing loaded yet, just switch voices and exit.
        if not has_audio_text:
            try:
                self.tts_manager.set_voice_model(model_name)
                self.current_language = selection
                self.progress_bar_manager.reset_progress_bar()
                self.update_play_pause_button(False)
                self.text_manager.hide_results()
            except FileNotFoundError as exc:
                messagebox.showerror("Voice Not Found", str(exc))
                revert_selection()
            except Exception as exc:
                messagebox.showerror("Voice Error", f"Could not switch voice:\n{exc}")
                revert_selection()
            return

        file_key = getattr(self, "current_file_key", None)
        new_lang_path = self.get_generation_path(file_key, language=selection) if file_key else None
        new_audio_exists = bool(new_lang_path and Path(new_lang_path).is_file())

        try:
            prompt_title = "Existing Audio Found" if new_audio_exists else "Regenerate Audio"
            if new_audio_exists:
                response = messagebox.askyesnocancel(
                    prompt_title,
                    f"Audio for this document already exists in {selection}.\n"
                    "Yes: Use existing audio\nNo: Regenerate in new language\nCancel: Keep current language/audio"
                )
                if response is None:
                    revert_selection()
                    return
                # Apply language change now
                self.tts_manager.set_voice_model(model_name)
                self.progress_bar_manager.reset_progress_bar()
                self.update_play_pause_button(False)
                self.text_manager.hide_results()
                self.current_language = selection
                if response is True:
                    # Use saved audio
                    if new_lang_path and not self.load_existing_generation(new_lang_path, existing_text):
                        self.regeneration_reason = "language"
                        self.generate_tts_in_background(existing_text, save_key=file_key, language=selection)
                else:
                    self.regeneration_reason = "language"
                    self.generate_tts_in_background(existing_text, save_key=file_key, language=selection)
            else:
                regenerate = messagebox.askyesno(
                    prompt_title,
                    f"Would you like to generate the audio in {selection}?"
                )
                if not regenerate:
                    revert_selection()
                    return
                self.tts_manager.set_voice_model(model_name)
                self.progress_bar_manager.reset_progress_bar()
                self.update_play_pause_button(False)
                self.text_manager.hide_results()
                self.current_language = selection
                self.regeneration_reason = "language"
                self.generate_tts_in_background(existing_text, save_key=file_key, language=selection)
            if self.current_is_admin:
                self.save_ui_settings()
        except FileNotFoundError as exc:
            messagebox.showerror("Voice Not Found", str(exc))
            revert_selection()
        except Exception as exc:
            messagebox.showerror("Voice Error", f"Could not switch voice:\n{exc}")
            revert_selection()

    def try_show_file_loaded_message(self):
        # Deprecated shim; route to unified message handler
        self.try_show_pending_messages()

    def try_show_pending_messages(self):
        if self.details_dialog_open:
            return

        if self.pending_audio_ready_message:
            self.pending_audio_ready_message = False
            messagebox.showinfo("Audio Ready", "New audio has been generated and loaded.")

        if self.pending_file_loaded_message and getattr(self, "tts_from_file", False):
            self.pending_file_loaded_message = False
            messagebox.showinfo("File Loaded", "The file has been successfully loaded for the typing test.")

    def load_saved_details(self, details_path):
        details_path = Path(details_path)
        if details_path.is_file():
            try:
                with open(details_path, "r", encoding="utf-8") as file:
                    data = json.load(file)
                return data.get("details", [])
            except Exception:
                return []
        return []

    def save_details(self, details_path, details):
        payload = {
            "details": details,
            "saved_at": time.time()
        }
        with open(details_path, "w", encoding="utf-8") as file:
            json.dump(payload, file, ensure_ascii=False, indent=2)

    def show_details_selection_dialog(self, text_content, initial_details):
        self.details_dialog_open = True
        dialog = tk.Toplevel(self.root)
        dialog.title("Select Important Details")
        dialog.grab_set()
        dialog.transient(self.root)
        self.bring_window_to_front(dialog)

        instructions = tk.Label(
            dialog,
            text="Highlight text below and click 'Add Detail' to track it during grading.",
            font=("Arial", 12)
        )
        instructions.pack(pady=(10, 5))

        text_frame = tk.Frame(dialog)
        text_frame.pack(fill="both", expand=True, padx=10, pady=5)

        text_widget = tk.Text(text_frame, wrap="word", height=15)
        text_widget.insert("1.0", text_content)
        text_widget.pack(side="left", fill="both", expand=True)
        text_widget.bind("<Key>", lambda event: "break")

        text_scroll = ttk.Scrollbar(text_frame, orient="vertical", command=text_widget.yview)
        text_scroll.pack(side="right", fill="y")
        text_widget.config(yscrollcommand=text_scroll.set)

        list_label = tk.Label(dialog, text="Selected Details:", font=("Arial", 11))
        list_label.pack(pady=(10, 0))

        listbox = tk.Listbox(dialog, height=6)
        listbox.pack(fill="x", padx=10)

        details = [detail for detail in initial_details if detail.strip()]
        for detail in details:
            listbox.insert("end", detail)

        button_frame = tk.Frame(dialog)
        button_frame.pack(fill="x", pady=10)

        def add_detail():
            try:
                selection = text_widget.get("sel.first", "sel.last").strip()
            except tk.TclError:
                selection = ""
            cleaned = " ".join(selection.split())
            if cleaned and cleaned not in details:
                details.append(cleaned)
                listbox.insert("end", cleaned)

        def remove_detail():
            selection = listbox.curselection()
            if selection:
                index = selection[0]
                detail = listbox.get(index)
                details.remove(detail)
                listbox.delete(index)

        def clear_all_details():
            if details:
                details.clear()
                listbox.delete(0, "end")

        result = {"value": None}

        def confirm():
            result["value"] = details.copy()
            dialog.destroy()

        def cancel():
            result["value"] = None
            dialog.destroy()

        add_button = tk.Button(button_frame, text="Add Detail", command=add_detail)
        add_button.pack(side="left", padx=5)

        remove_button = tk.Button(button_frame, text="Remove Detail", command=remove_detail)
        remove_button.pack(side="left", padx=5)

        clear_button = tk.Button(button_frame, text="Clear All", command=clear_all_details)
        clear_button.pack(side="left", padx=5)

        spacer = tk.Frame(button_frame)
        spacer.pack(side="left", expand=True)

        save_button = tk.Button(button_frame, text="Save Details", command=confirm)
        save_button.pack(side="right", padx=5)

        cancel_button = tk.Button(button_frame, text="Cancel", command=cancel)
        cancel_button.pack(side="right", padx=5)

        dialog.protocol("WM_DELETE_WINDOW", cancel)
        try:
            self.fit_window_to_content(dialog, min_size=(800, 600))
            dialog.wait_window()
        finally:
            self.details_dialog_open = False
        self.try_show_pending_messages()
        return result["value"]

    def build_road_variant_map(self):
        mapping = {}
        for canonical, variants in ROAD_VARIATIONS.items():
            for variant in variants:
                mapping[variant.lower()] = canonical
        return mapping

    def apply_saved_settings(self):
        settings = self.load_ui_settings()
        if not settings:
            self.update_admin_controls()
            return

        distortion = settings.get("distortion")
        if distortion in ("on_distortion", "off_distortion"):
            self.distortion_status.set(distortion)
            self.update_distortion_setting(force=True)

        language = settings.get("language")
        if language in self.voice_options:
            self.language_var.set(language)
            self.change_language(force=True)

        try:
            speed = float(settings.get("speed", 0))
            if speed:
                self.speed_var.set(speed)
        except Exception:
            pass
        self.speed_dirty = False
        self.update_apply_speed_button()

        highlight = settings.get("highlight")
        if highlight in ("on_highlight", "off_highlight"):
            self.highlight_var.set(highlight)

        self.update_admin_controls()

    def normalize_words(self, text):
        cleaned = re.sub(r'[^\w\s]', ' ', text).lower().split()
        return [self.road_variant_map.get(word, word) for word in cleaned]

    def normalize_text_for_matching(self, text):
        return " ".join(self.normalize_words(text))

    def submit_text(self):
        self.stop_timer_display()
        user_text = self.text_manager.get_text()
        reference = self.tts_manager.getTypingText()
        word_count = len(user_text.split())
        elapsed_time = time.time() - self.start_time if self.start_time else 1
        wpm = word_count / (elapsed_time / 60) if elapsed_time > 0 else 0

        accuracy = self.calculate_word_accuracy(user_text, reference)
        details_score = self.calculate_details_score(user_text)
        details_text = f"{details_score:.2f}" if details_score is not None else "N/A"

        results = (
            f"You typed {word_count} words.\n"
            f"Words per Minute: {wpm:.2f}\n"
            f"Accuracy: {accuracy:.2f}\n"
            f"Details: {details_text}"
        )
        self.progress_bar_manager.hide_progress_bar()
        self.text_manager.show_results(results)
        self.text_manager.highlight_submission_errors(self.tts_manager.getTypingText())
        username = self.current_username if self.current_username else "Guest"
        self.save_score_to_csv(username, wpm, accuracy, details_score)
        messagebox.showinfo("Score Saved", f"Results saved for {username}.")
        self.root.after(5000, self.reset_ui) 

    def reset_ui(self):
        self.reset_for_new_audio()

    def discard_text(self):
        self.stop_timer_display()
        self.text_manager.clear_text()
        self.start_time = None

    def update_play_pause_button(self, playing=False):
        if hasattr(self, "play_pause_button"):
            self.play_pause_button.config(text="⏸" if playing else "▶")

    def is_audio_playing(self):
        stream = getattr(self.tts_manager, "stream", None)
        active = bool(getattr(stream, "active", False))
        return stream is not None and active and not self.tts_manager.is_paused

    def is_audio_paused(self):
        stream = getattr(self.tts_manager, "stream", None)
        return stream is not None and self.tts_manager.is_paused

    def handle_playback_complete(self):
        # Sync button state when audio naturally ends
        self.tts_manager.pauseTTS()
        self.update_play_pause_button(False)

    def pause_audio(self):
        self.tts_manager.pauseTTS()
        self.progress_bar_manager.pause_progress_bar()
        self.update_play_pause_button(False)

    def reset_for_new_audio(self):
        self.stop_timer_display()
        self.start_time = None
        self.text_manager.hide_results()
        # Clear any previous highlights/errors from the typing box
        self.text_manager.typing_box.tag_remove("error", "1.0", "end")
        self.text_manager.typing_box.tag_remove("correct", "1.0", "end")
        self.text_manager.typing_box.tag_remove("incorrect", "1.0", "end")
        self.text_manager.clear_text()
        self.reset_audio(auto_resume=False)
        self.update_apply_speed_button()

    def toggle_play_pause(self):
        if not self.tts_manager.getTypingText().strip():
            messagebox.showwarning("No Document Loaded", "Please load a document before starting the audio.")
            return

        paused = self.is_audio_paused()
        speed = self.speed_var.get()
        need_prepare = not self.tts_manager.is_armed

        if self.is_audio_playing():
            self.pause_audio()
            return

        if getattr(self.tts_manager, "use_synth_speed", False):
            try:
                target_scale = self.tts_manager._to_piper_scale(speed)
                last_scale = getattr(self.tts_manager, "_last_synth_scale", None)
                if last_scale is None or abs(target_scale - last_scale) > 1e-6:
                    need_prepare = True
            except Exception:
                need_prepare = True

        if getattr(self.tts_manager, "playback_finished", False):
            self.tts_manager.reset_playback()
            paused = False
            need_prepare = True

        if paused and not need_prepare:
            self.tts_manager.resumeTTS()
            self.progress_bar_manager.resume_progress_bar()
        else:
            if need_prepare:
                self.tts_manager.prepareTTS(speed=speed)
            self.tts_manager.playTTS(speed=speed)
            self.progress_bar_manager.start_progress_bar(speed=speed)

        self.update_play_pause_button(True)

    def reset_audio(self, auto_resume=None):
        """Reset playback to start; optionally resume if it was playing."""
        was_playing = self.is_audio_playing()
        should_resume = was_playing if auto_resume is None else bool(auto_resume)
        has_text = bool(self.tts_manager.getTypingText().strip())
        speed = self.speed_var.get()

        self.tts_manager.reset_playback()
        self.progress_bar_manager.reset_progress_bar()
        if self.tts_manager.audio_data is not None:
            self.progress_bar_manager.update_audio_duration(speed=speed)

        if not has_text or self.tts_manager.audio_data is None:
            self.update_play_pause_button(False)
            return

        # Re-arm audio at the beginning
        self.tts_manager.prepareTTS(speed=speed)

        if should_resume:
            self.tts_manager.playTTS(speed=speed)
            self.progress_bar_manager.start_progress_bar(speed=speed)
            self.update_play_pause_button(True)
        else:
            self.update_play_pause_button(False)

    def on_speed_dirty(self, *_):
        # Mark speed as needing apply and update button label/state
        self.speed_dirty = True
        self.update_apply_speed_button()

    def update_apply_speed_button(self):
        if not hasattr(self, "apply_speed_button"):
            return
        label = f"Apply Speed ({self.speed_var.get():.1f}x)"
        self.apply_speed_button.config(text=label)
        manager = getattr(self, "tts_manager", None)
        text_loaded = bool(manager and manager.getTypingText().strip())
        enabled = self.speed_dirty and text_loaded and not self.generating and self.current_is_admin
        self.apply_speed_button.config(state="normal" if enabled else "disabled")

    def apply_speed_change(self):
        if not self.current_is_admin:
            messagebox.showwarning("Admin Only", "Speed changes are available to admins only.")
            return
        if self.generating:
            return

        text = self.tts_manager.getTypingText()
        if not text.strip():
            messagebox.showwarning("No Document Loaded", "Please load a document before applying speed changes.")
            self.speed_dirty = False
            self.update_apply_speed_button()
            return

        # Immediately stop and reset playback/test state
        self.reset_for_new_audio()
        self.speed_dirty = False
        self.update_apply_speed_button()

        self.generate_tts_in_background(
            text,
            save_key=getattr(self, "current_file_key", None),
            language=self.current_language,
            message="Regenerating audio..."
        )
        self.regeneration_reason = "speed"
        self.save_ui_settings()

    def show_loading_window(self, message="Generating TTS..."):
        self.generating = True
        self.update_apply_speed_button()
        self.loading_window = tk.Toplevel(self.root)
        self.loading_window.title("Please Wait")
        self.loading_window.resizable(False, False)
        self.loading_window.grab_set()
        self.loading_window.transient(self.root)
        self.bring_window_to_front(self.loading_window)

        label = tk.Label(self.loading_window, text=message, font=("Arial", 12))
        label.pack(pady=10)

        progress = ttk.Progressbar(self.loading_window, mode='indeterminate', length=200)
        progress.pack(pady=10)
        progress.start()
        
        self.fit_window_to_content(self.loading_window, min_size=(320, 140))

    def hide_loading_window(self):
        if hasattr(self, "loading_window") and self.loading_window.winfo_exists():
            self.loading_window.destroy()
        self.generating = False
        self.update_apply_speed_button()

    def on_tts_ready(self, reload_path=None, text_content=None):
        self.hide_loading_window()
        if reload_path and Path(reload_path).is_file():
            self.load_existing_generation(reload_path, text_content or self.tts_manager.getTypingText(), show_message=False)
        self.progress_bar_manager.update_audio_duration(speed=self.speed_var.get())
        self.reset_for_new_audio()
        self.speed_dirty = False
        self.update_apply_speed_button()
        self.pending_audio_ready_message = True

        if getattr(self, "tts_from_file", False) and not self.regeneration_reason:
            self.pending_file_loaded_message = True
        self.try_show_pending_messages()
        self.regeneration_reason = None

    def generate_tts_in_background(self, text, save_key=None, language=None, message="Generating TTS..."):
        def task():
            self.tts_manager.TTSGenerate(text)
            self.tts_manager.prepareTTS(speed=self.speed_var.get())
            reload_path = None
            if save_key:
                self.save_generation_copy(save_key, language=language)
                reload_path = self.get_generation_path(save_key, language=language)
            self.root.after(0, lambda: self.on_tts_ready(reload_path, text))

        self.show_loading_window(message)
        threading.Thread(target=task, daemon=True).start()

    def save_generation_copy(self, file_key, language=None):
        try:
            target = self.get_generation_path(file_key, language=language)
            shutil.copy2(self.tts_manager.wav_file, target)
        except Exception:
            pass

    def on_typing(self, event):
        user_input = self.text_manager.get_text()
        reference = self.tts_manager.getTypingText()
        highlight_enabled = self.highlight_var.get() == "on_highlight"
        self.text_manager.highlight_typing_progress(user_input, reference, highlight_enabled)

    def calculate_word_accuracy(self, user_text, reference_text):
        user_words = self.normalize_words(user_text)
        reference_words = self.normalize_words(reference_text)

        # Word-level Levenshtein distance
        def levenshtein(seq1, seq2):
            m, n = len(seq1), len(seq2)
            dp = [[0] * (n+1) for _ in range(m+1)]

            for i in range(m+1):
                dp[i][0] = i
            for j in range(n+1):
                dp[0][j] = j

            for i in range(1, m+1):
                for j in range(1, n+1):
                    cost = 0 if seq1[i-1] == seq2[j-1] else 1
                    dp[i][j] = min(
                        dp[i-1][j] + 1,     # deletion
                        dp[i][j-1] + 1,     # insertion
                        dp[i-1][j-1] + cost # substitution
                    )
            return dp[m][n]

        distance = levenshtein(user_words, reference_words)
        total = max(len(reference_words), len(user_words), 1)  # avoid div by zero
        accuracy = (1 - distance / total) * 100
        return accuracy

    def calculate_details_score(self, user_text):
        details = [detail for detail in self.current_details if detail.strip()]
        if not details:
            return None

        user_normalized = self.normalize_text_for_matching(user_text)
        matches = 0
        for detail in details:
            normalized_detail = self.normalize_text_for_matching(detail)
            if normalized_detail and normalized_detail in user_normalized:
                matches += 1
        return (matches / len(details)) * 100
    
    def save_score_to_csv(self, username, wpm, accuracy, details_score):
        from datetime import datetime
        store = self.load_scores_store() or {}

        first_name = self.current_first_name or ("Guest" if username == "Guest" else "N/A")
        last_name = self.current_last_name or ("User" if username == "Guest" else "N/A")
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        entry = {
            "test_no": len(store.get(username, [])) + 1,
            "time": timestamp,
            "wpm": f"{wpm:.2f}",
            "accuracy": f"{accuracy:.2f}",
            "details": f"{details_score:.2f}" if details_score is not None else "N/A",
            "first_name": first_name,
            "last_name": last_name
        }

        store.setdefault(username, []).append(entry)
        self.save_scores_store(store)

    def load_scores_store(self):
        if not self.scores_file.exists():
            return {}
        try:
            with open(self.scores_file, "rb") as f:
                payload = f.read()
            return self.decrypt_payload(payload).get("scores", {})
        except Exception:
            return {}

    def save_scores_store(self, store: dict):
        payload = {"scores": store}
        cipher = self.encrypt_payload(payload)
        self.scores_file.parent.mkdir(parents=True, exist_ok=True)
        with open(self.scores_file, "wb") as f:
            f.write(cipher)

    def on_highlight_changed(self):
        # Only admins can toggle; save preference when allowed
        if not self.current_is_admin:
            return
        self.save_ui_settings()


    
    def start_timer_display(self):
        if self.timer_id is None:
            self.text_manager.timer_label.grid()
            self.update_timer_display()

    def update_timer_display(self):
        if self.start_time:
            elapsed = time.time() - self.start_time
            self.text_manager.timer_label.config(text=f"Time: {elapsed:.1f}s")
            self.timer_id = self.root.after(500, self.update_timer_display)

    def stop_timer_display(self):
        if self.timer_id:
            self.root.after_cancel(self.timer_id)
            self.timer_id = None
        self.text_manager.timer_label.grid_remove()


    def start_timer_if_needed(self, event=None):
        if self.start_time is None:
            self.start_time = time.time()
            self.start_timer_display()
